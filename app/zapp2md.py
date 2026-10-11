#!/usr/bin/env python3
"""
zapp2md - Converts a WhatsApp chat export (with media) into a single chronological
Markdown file, ready to be analysed with an LLM such as Claude, ChatGPT or Gemini, or other AI tools. Interface and output: Italiano / Português / English.

- Audio / voice messages: full local transcription (faster-whisper).
- Video: transcription of the audio track (optional).
- Images: OCR with the built-in Windows engine + EXIF metadata.
- Documents: text of PDF (with OCR for scanned pages), DOCX, XLSX, PPTX, TXT/CSV/VCF...

Everything runs locally, at no cost. Results are cached: if processing is interrupted,
running it again resumes where it stopped.
"""

from __future__ import annotations

import os

# Download "classico" dei modelli: il protocollo xet di Hugging Face può chiudere il processo su Windows.
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")
os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
os.environ.setdefault("HF_HUB_DISABLE_TELEMETRY", "1")
os.environ.setdefault("HF_HUB_VERBOSITY", "error")
# Versione installata: il modello di trascrizione sta nella cartella "models" del programma
# (così la disinstallazione lo rimuove). Senza quella cartella si usa la cache standard dell'utente.
_MODELS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "models")
if os.path.isdir(_MODELS_DIR):
    os.environ.setdefault("HF_HUB_CACHE", _MODELS_DIR)

# IMPORTANTE: il motore di trascrizione (ctranslate2) deve essere caricato PRIMA delle librerie
# WinRT dell'OCR, altrimenti su Windows il caricamento del modello termina con un crash nativo.
try:
    import ctranslate2  # noqa: F401
    import faster_whisper  # noqa: F401
except Exception:
    pass

import argparse
import asyncio
import datetime as dt
import json
import logging
import re
import sqlite3
import sys
import time
import traceback
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from i18n import LANGS, get_lang, set_lang, t  # noqa: E402

VERSION = "1.0"
AUTHOR = "Andrea Moretto"
AUTHOR_URL = "https://www.linkedin.com/in/andrea-moretto-mba/"
CACHE_SCHEMA = "1"
DEFAULT_MODEL = "large-v3-turbo"

# Eventi di avanzamento in formato JSON (una riga per evento) per l'interfaccia grafica.
_JSON_EVENTS = False


def emit(ev: str, **kw) -> None:
    if _JSON_EVENTS:
        sys.stdout.write(json.dumps({"ev": ev, **kw}, ensure_ascii=False) + "\n")
        sys.stdout.flush()


class JsonLogHandler(logging.Handler):
    def emit(self, record):  # noqa: D401 - firma imposta da logging
        try:
            emit("log", level=record.levelname, msg=record.getMessage())
        except Exception:
            pass

# --------------------------------------------------------------------------- #
# Costanti
# --------------------------------------------------------------------------- #

AUDIO_EXT = {".opus", ".ogg", ".oga", ".m4a", ".mp3", ".aac", ".wav", ".amr", ".flac", ".wma", ".weba"}
VIDEO_EXT = {".mp4", ".mov", ".3gp", ".mkv", ".avi", ".webm", ".m4v", ".mpeg", ".mpg", ".wmv"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic", ".heif"}
DOC_EXT = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}
TEXT_EXT = {".txt", ".csv", ".tsv", ".vcf", ".json", ".xml", ".html", ".htm", ".md", ".log", ".ics"}

INVISIBLE = dict.fromkeys(map(ord, "\u200e\u200f\u202a\u202b\u202c\u202d\u202e\u2066\u2067\u2068\u2069\ufeff"), None)

DATE_RE = r"\d{1,4}[./-]\d{1,2}[./-]\d{1,4}"
TIME_RE = r"\d{1,2}[:.]\d{2}(?:[:.]\d{2})?(?:\s?[AaPp]\.?\s?[Mm]\.?)?"
HEADER_RE = re.compile(
    rf"^(?:\[(?P<d1>{DATE_RE}),?\s+(?P<t1>{TIME_RE})\]\s?|(?P<d2>{DATE_RE}),?\s+(?P<t2>{TIME_RE})\s[-\u2013]\s)(?P<rest>.*)$"
)
# Android: "IMG-20240115-WA0003.jpg (file allegato)" / "(file attached)" / "(ficheiro anexado)"...
ANDROID_ATTACH_RE = re.compile(r"^(?P<f>.+?\.[A-Za-z0-9]{1,6})\s\((?P<lbl>[^()]{2,40})\)\s*$")
ATTACH_LABEL_RE = re.compile(r"allegat|attach|anex|adjunt|joint|angeh", re.I)
# iOS: "<allegato: 00000012-PHOTO-2024-01-15-14-05-33.jpg>" / "<attached: ...>" / "<anexado: ...>"
IOS_ATTACH_RE = re.compile(r"<(?P<lbl>[^<>:]{2,30}):\s*(?P<f>[^<>]+?\.[A-Za-z0-9]{1,6})>")
# Media esclusi dall'esportazione ("Esporta senza media"): "<messaggio vocale omesso>", "<Media omitted>",
# "audio omitted", "<Multimédia oculta>", "<documento omitido> nome.pdf"...
OMITTED_RE = re.compile(r"<[^<>]*(?:omess|omitted|omitid|ocult)[^<>]*>|\b(?:audio|image|video|sticker|document|GIF)\s+omitted\b",
                        re.I)
OMITTED_VOICE_RE = re.compile(r"vocal|voice|voz|audio|áudio", re.I)


def count_omitted(msgs) -> tuple[int, int]:
    """(allegati omessi, di cui vocali/audio)."""
    hits = [mo.group(0) for msg in msgs for mo in OMITTED_RE.finditer(msg.text)]
    return len(hits), sum(1 for h in hits if OMITTED_VOICE_RE.search(h))


FILENAME_DATE_RES = [
    re.compile(r"(?:IMG|VID|AUD|PTT|DOC|STK|GIF)-(\d{4})(\d{2})(\d{2})-WA\d+", re.I),
    re.compile(r"-(\d{4})-(\d{2})-(\d{2})-\d{2}-\d{2}-\d{2}"),
]
TITLE_PREFIX_RE = re.compile(
    r"^(Chat WhatsApp con|Chat di WhatsApp con|WhatsApp Chat with|WhatsApp Chat -|WhatsApp Chat|"
    r"Conversa do WhatsApp com|Conversa no WhatsApp com|Conversa de WhatsApp com|Conversa WhatsApp com)\s*", re.I)

log = logging.getLogger("zapp2md")


# --------------------------------------------------------------------------- #
# Utilità
# --------------------------------------------------------------------------- #

def clean(s: str) -> str:
    return s.translate(INVISIBLE).replace("\u202f", " ").replace("\u00a0", " ")


def fmt_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return str(n)


def fmt_dur(sec: float | None) -> str:
    if sec is None:
        return "?"
    sec = int(round(sec))
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def md_escape_line(line: str) -> str:
    """Evita che il testo dei messaggi venga scambiato per struttura Markdown."""
    if line.startswith(("#", ">", "---", "===", "```")):
        return "\\" + line
    return line


def quote(text: str) -> list[str]:
    return ["> " + l if l.strip() else ">" for l in text.splitlines()] or [">"]


def date_from_filename(name: str) -> str | None:
    for rx in FILENAME_DATE_RES:
        m = rx.search(name)
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
            except ValueError:
                return None
    return None


def media_kind(name: str) -> str:
    ext = Path(name).suffix.lower()
    up = name.upper()
    if ext in AUDIO_EXT:
        return "audio"
    if ext in VIDEO_EXT:
        return "video"
    if ext in IMAGE_EXT:
        return "sticker" if up.startswith("STK-") or "-STICKER-" in up else "image"
    if ext in DOC_EXT or ext in TEXT_EXT:
        return "document"
    return "other"


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


class UserError(Exception):
    pass


# --------------------------------------------------------------------------- #
# Parsing della chat
# --------------------------------------------------------------------------- #

@dataclass
class Message:
    ts: dt.datetime
    sender: str | None  # None = messaggio di sistema
    text: str
    attachments: list[str] = field(default_factory=list)  # nomi file (presenti o mancanti)


def detect_date_order(dates: list[str]) -> str:
    firsts, seconds = [], []
    for d in dates:
        a, b, _ = re.split(r"[./-]", d)
        if len(a) == 4:
            return "ymd"
        firsts.append(int(a))
        seconds.append(int(b))
    if any(x > 12 for x in firsts):
        return "dmy"
    if any(x > 12 for x in seconds):
        return "mdy"
    return "dmy"


def parse_ts(d: str, tm: str, order: str) -> dt.datetime:
    p = [int(x) for x in re.split(r"[./-]", d)]
    if order == "ymd":
        y, mo, da = p
    elif order == "mdy":
        mo, da, y = p
    else:
        da, mo, y = p
    if y < 100:
        y += 2000
    ampm = None
    m = re.search(r"([AaPp])\.?\s?[Mm]\.?$", tm)
    if m:
        ampm = m.group(1).lower()
        tm = tm[: m.start()].strip()
    tp = [int(x) for x in re.split(r"[:.]", tm)]
    hh, mm = tp[0], tp[1]
    ss = tp[2] if len(tp) > 2 else 0
    if ampm == "p" and hh < 12:
        hh += 12
    elif ampm == "a" and hh == 12:
        hh = 0
    return dt.datetime(y, mo, da, hh, mm, ss)


def count_headers(path: Path, max_lines: int = 300) -> int:
    try:
        lines = read_text_file(path).splitlines()[:max_lines]
    except Exception:
        return 0
    return sum(1 for l in lines if HEADER_RE.match(clean(l)))


def find_chat_file(files: dict[str, Path]) -> Path:
    txts = [p for p in files.values() if p.suffix.lower() == ".txt"]
    if not txts:
        raise UserError(t("err_no_txt"))
    scored = sorted(((count_headers(p), p) for p in txts), key=lambda x: (-x[0], len(x[1].name)))
    if scored[0][0] == 0:
        raise UserError(t("err_format", f=", ".join(p.name for p in txts[:5])))
    return scored[0][1]


def parse_chat(path: Path, media: dict[str, Path], date_order: str) -> tuple[list[Message], str]:
    raw: list[list] = []  # [data, ora, resto, righe_successive]
    for line in read_text_file(path).splitlines():
        cl = clean(line)
        m = HEADER_RE.match(cl)
        if m:
            raw.append([m.group("d1") or m.group("d2"), m.group("t1") or m.group("t2"), m.group("rest"), []])
        elif raw:
            raw[-1][3].append(cl)
    if date_order == "auto":
        date_order = detect_date_order([r[0] for r in raw]) if raw else "dmy"

    media_lower = {k.lower(): k for k in media}

    def resolve(name: str) -> str | None:
        name = name.strip()
        return name if name in media else media_lower.get(name.lower())

    msgs: list[Message] = []
    for d, tm, rest, extra in raw:
        try:
            ts = parse_ts(d, tm, date_order)
        except ValueError:
            # riga con data non valida: la tratto come testo del messaggio precedente
            if msgs:
                msgs[-1].text += "\n" + f"{d} {tm} {rest}"
            continue
        sender, text = None, rest
        idx = rest.find(": ")
        if 0 < idx <= 80:
            sender, text = rest[:idx].strip(), rest[idx + 2:]
        body = "\n".join([text] + extra)

        attachments: list[str] = []
        kept: list[str] = []
        for bl in body.split("\n"):
            found = False
            mi = IOS_ATTACH_RE.search(bl)
            if mi:
                fn = mi.group("f").strip()
                attachments.append(resolve(fn) or fn)
                bl = (bl[: mi.start()] + bl[mi.end():]).strip()
                found = True
            else:
                ma = ANDROID_ATTACH_RE.match(bl.strip())
                if ma:
                    r = resolve(ma.group("f"))
                    if r or ATTACH_LABEL_RE.search(ma.group("lbl")):
                        attachments.append(r or ma.group("f").strip())
                        bl, found = "", True
                if not found and bl.strip() and " " not in bl.strip():
                    r = resolve(bl.strip())
                    if r:
                        attachments.append(r)
                        bl, found = "", True
            if bl or not found:
                kept.append(bl)
        msgs.append(Message(ts, sender, "\n".join(kept).strip("\n"), attachments))
    return msgs, date_order


# --------------------------------------------------------------------------- #
# Estrattori
# --------------------------------------------------------------------------- #

class WinOCR:
    """OCR con il motore integrato di Windows 10/11 (gratuito, offline)."""

    def __init__(self, lang: str):
        from winrt.windows.globalization import Language
        from winrt.windows.graphics.imaging import BitmapPixelFormat, SoftwareBitmap
        from winrt.windows.media.ocr import OcrEngine
        from winrt.windows.storage.streams import DataWriter

        self._SoftwareBitmap, self._Fmt, self._DataWriter = SoftwareBitmap, BitmapPixelFormat, DataWriter
        self.available = [l.language_tag for l in OcrEngine.available_recognizer_languages]
        engine = None
        if lang and lang != "auto":
            for tag in self.available:
                if tag.lower().split("-")[0] == lang.lower().split("-")[0]:
                    engine = OcrEngine.try_create_from_language(Language(tag))
                    self.lang = tag
                    break
        if engine is None:
            engine = OcrEngine.try_create_from_user_profile_languages()
            self.lang = "user profile"
        if engine is None:
            raise RuntimeError("no OCR language installed in Windows")
        self.engine = engine
        self.max_dim = int(OcrEngine.max_image_dimension)
        self.loop = asyncio.new_event_loop()

    def recognize(self, img) -> str:
        from PIL import Image

        img = img.convert("RGBA")
        w, h = img.size
        if max(w, h) > self.max_dim:
            s = self.max_dim / max(w, h)
            img = img.resize((max(1, int(w * s)), max(1, int(h * s))), Image.LANCZOS)
        elif max(w, h) < 800:  # le immagini piccole si leggono meglio ingrandite
            s = min(3.0, 1600 / max(w, h, 1))
            img = img.resize((max(1, int(w * s)), max(1, int(h * s))), Image.LANCZOS)
        w, h = img.size
        dw = self._DataWriter()
        dw.write_bytes(img.tobytes("raw", "BGRA"))
        sb = self._SoftwareBitmap.create_copy_from_buffer(dw.detach_buffer(), self._Fmt.BGRA8, w, h)
        result = self.loop.run_until_complete(self.engine.recognize_async(sb))
        return "\n".join(line.text for line in result.lines).strip()


class Transcriber:
    def __init__(self, model_name: str, lang: str, threads: int):
        from faster_whisper import WhisperModel

        log.info(t("loading_model", m=model_name))
        self.model = WhisperModel(model_name, device="cpu", compute_type="int8", cpu_threads=threads)
        self.lang = None if lang == "auto" else lang

    def transcribe(self, path: Path) -> dict:
        opts = dict(beam_size=5, vad_filter=True, condition_on_previous_text=False,
                    vad_parameters=dict(min_silence_duration_ms=700))
        segs, info = self.model.transcribe(str(path), language=None, **opts)
        lang, prob = info.language, info.language_probability
        # Sui vocali brevi il riconoscimento automatico della lingua può sbagliare:
        # se è incerto si usa la lingua principale della chat.
        if self.lang and lang != self.lang and prob < 0.80:
            segs, info = self.model.transcribe(str(path), language=self.lang, **opts)
            lang = self.lang
        out = []
        for s in segs:
            txt = s.text.strip()
            if txt:
                out.append([round(s.start, 1), txt])
        return {"duration": info.duration, "language": lang, "segments": out}


def probe_av(path: Path) -> dict:
    import av

    meta: dict = {}
    with av.open(str(path)) as c:
        if c.duration:
            meta["duration"] = c.duration / 1_000_000
        v = next((s for s in c.streams if s.type == "video"), None)
        a = next((s for s in c.streams if s.type == "audio"), None)
        if v is not None and v.codec_context.width:
            meta["resolution"] = f"{v.codec_context.width}×{v.codec_context.height}"
        meta["has_audio"] = a is not None
        ct = (c.metadata or {}).get("creation_time")
        if ct and not ct.startswith("1970"):
            meta["recorded"] = ct[:19].replace("T", " ")
    return meta


def exif_info(img) -> dict:
    out: dict = {}
    try:
        ex = img.getexif()
    except Exception:
        return out
    if not ex:
        return out
    try:
        sub = ex.get_ifd(0x8769)
    except Exception:
        sub = {}
    date = sub.get(36867) or ex.get(306)
    if date:
        out["taken"] = str(date).strip().replace(":", "-", 2)
    cam = " ".join(str(x).strip() for x in (ex.get(271), ex.get(272)) if x)
    if cam:
        out["camera"] = cam
    try:
        gps = ex.get_ifd(0x8825)
        if gps and 2 in gps and 4 in gps:
            def dec(v, ref):
                deg = float(v[0]) + float(v[1]) / 60 + float(v[2]) / 3600
                return -deg if ref in ("S", "W") else deg
            out["gps"] = f"{dec(gps[2], gps.get(1, 'N')):.5f}, {dec(gps[4], gps.get(3, 'E')):.5f}"
    except Exception:
        pass
    return out


class Extractor:
    def __init__(self, args):
        self.args = args
        self._ocr = None
        self._ocr_failed = False
        self._tr = None

    @property
    def ocr(self) -> WinOCR | None:
        if self._ocr is None and not self._ocr_failed and not self.args.no_ocr:
            try:
                self._ocr = WinOCR(self.args.speech_lang)
                log.info(t("ocr_on", l=self._ocr.lang, a=", ".join(self._ocr.available)))
            except Exception as e:
                self._ocr_failed = True
                log.warning(t("ocr_off", e=e))
        return self._ocr

    @property
    def transcriber(self) -> Transcriber:
        if self._tr is None:
            self._tr = Transcriber(self.args.model, self.args.speech_lang, self.args.threads)
        return self._tr

    # ---- immagini --------------------------------------------------------- #
    def image(self, path: Path) -> dict:
        from PIL import Image, ImageOps

        if path.suffix.lower() in (".heic", ".heif"):
            try:
                import pillow_heif  # type: ignore
                pillow_heif.register_heif_opener()
            except ImportError:
                return {"meta": {}, "note": "n_heic"}
        with Image.open(path) as im:
            meta: dict = {"dimensions": f"{im.width}×{im.height}"}
            meta.update(exif_info(im))
            if getattr(im, "n_frames", 1) > 1:
                meta["animated"] = True
                im.seek(0)
            frame = ImageOps.exif_transpose(im.copy())
        text = self.ocr.recognize(frame) if self.ocr else ""
        return {"meta": meta, "text": text}

    # ---- audio / video ---------------------------------------------------- #
    def audio(self, path: Path) -> dict:
        try:
            meta = probe_av(path)
        except Exception:
            meta = {}
        meta.pop("has_audio", None)
        meta.pop("resolution", None)
        r = self.transcriber.transcribe(path)
        meta["duration"] = r["duration"] or meta.get("duration")
        meta["language"] = r["language"]
        return {"meta": meta, "segments": r["segments"]}

    def video(self, path: Path) -> dict:
        meta = probe_av(path)
        has_audio = meta.pop("has_audio", False)
        if self.args.no_video:
            return {"meta": meta, "note": "n_video_off"}
        if not has_audio:
            return {"meta": meta, "note": "n_no_audio_track"}
        r = self.transcriber.transcribe(path)
        meta["language"] = r["language"]
        return {"meta": meta, "segments": r["segments"]}

    # ---- documenti -------------------------------------------------------- #
    def document(self, path: Path) -> dict:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._pdf(path)
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for ti, table in enumerate(d.tables, 1):
                parts.append(t("c_table", v=ti))
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            cp = d.core_properties
            meta = {k: v for k, v in (("title", cp.title), ("author", cp.author),
                                       ("created", str(cp.created)[:10] if cp.created else None)) if v}
            return {"meta": meta, "text": "\n".join(parts)}
        if ext in (".xlsx", ".xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            try:
                for ws in wb.worksheets:
                    parts.append(t("c_sheet", v=ws.title))
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i >= 5000:
                            parts.append(t("c_rows_omitted"))
                            break
                        vals = ["" if v is None else str(v) for v in row]
                        if any(vals):
                            parts.append("\t".join(vals).rstrip())
                n_sheets = len(wb.sheetnames)
            finally:
                wb.close()
            return {"meta": {"sheets": n_sheets}, "text": "\n".join(parts)}
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(t("c_slide", v=i))
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        parts.append(shape.text_frame.text)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    parts.append(t("c_notes") + slide.notes_slide.notes_text_frame.text)
            return {"meta": {"slides": len(prs.slides)}, "text": "\n".join(parts)}
        # file di testo
        text = read_text_file(path)
        if ext in (".html", ".htm"):
            text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"[ \t]+", " ", text)
        return {"meta": {}, "text": text.strip()}

    def _pdf(self, path: Path) -> dict:
        import pypdfium2 as pdfium

        try:
            doc = pdfium.PdfDocument(str(path))
        except pdfium.PdfiumError as e:
            if "password" in str(e).lower():
                return {"meta": {}, "note": "n_pdf_password"}
            raise
        try:
            meta: dict = {"pages": len(doc)}
            try:
                md = doc.get_metadata_dict(skip_empty=True)
            except Exception:
                md = {}
            for src, dst in (("Title", "title"), ("Author", "author")):
                if md.get(src):
                    meta[dst] = md[src]
            c = (md.get("CreationDate") or "").replace("D:", "")
            if len(c) >= 8 and c[:8].isdigit():
                meta["created"] = f"{c[:4]}-{c[4:6]}-{c[6:8]}"
            parts, ocr_pages = [], 0
            for i in range(len(doc)):
                page = doc[i]
                try:
                    tp = page.get_textpage()
                    txt = tp.get_text_range().replace("\r\n", "\n").replace("\r", "\n").strip()
                    tp.close()
                    if len(txt) < 30 and self.ocr and ocr_pages < self.args.max_ocr_pages:
                        img = page.render(scale=200 / 72).to_pil()  # 200 dpi
                        t2 = self.ocr.recognize(img)
                        ocr_pages += 1
                        if len(t2) > len(txt):
                            txt = t2
                finally:
                    page.close()
                if txt:
                    parts.append(t("c_page", v=i + 1) + "\n" + txt)
            if ocr_pages:
                meta["ocr_pages"] = ocr_pages
        finally:
            doc.close()
        return {"meta": meta, "text": "\n".join(parts)}

    # ---- dispatch --------------------------------------------------------- #
    def run(self, kind: str, path: Path) -> dict:
        if kind in ("image", "sticker"):
            return self.image(path)
        if kind == "audio":
            return self.audio(path)
        if kind == "video":
            return self.video(path)
        if kind == "document":
            return self.document(path)
        return {"meta": {}}


# --------------------------------------------------------------------------- #
# Cache (permette di riprendere un'elaborazione interrotta)
# --------------------------------------------------------------------------- #

class Cache:
    def __init__(self, path: Path):
        self.db = sqlite3.connect(str(path))
        self.db.execute("CREATE TABLE IF NOT EXISTS r (k TEXT PRIMARY KEY, v TEXT)")
        self.db.commit()

    def get(self, k: str):
        row = self.db.execute("SELECT v FROM r WHERE k=?", (k,)).fetchone()
        return json.loads(row[0]) if row else None

    def put(self, k: str, v) -> None:
        self.db.execute("INSERT OR REPLACE INTO r VALUES (?,?)", (k, json.dumps(v, ensure_ascii=False)))
        self.db.commit()

    def close(self):
        self.db.close()


def cache_key(args, name: str, path: Path, kind: str) -> str:
    sig = [CACHE_SCHEMA, kind, name, str(path.stat().st_size)]
    if kind in ("audio", "video"):
        sig += [args.model, args.speech_lang, "novideo" if (kind == "video" and args.no_video) else ""]
    if kind in ("image", "sticker", "document"):
        sig += ["noocr" if args.no_ocr else args.speech_lang, get_lang()]
    return "|".join(sig)


# --------------------------------------------------------------------------- #
# Input: zip o cartella
# --------------------------------------------------------------------------- #

def prepare_input(inp: Path, workdir: Path) -> Path:
    if inp.is_dir():
        return inp
    if inp.suffix.lower() == ".txt":
        return inp.parent
    if inp.suffix.lower() != ".zip":
        raise UserError(t("err_input", p=inp))
    done = workdir / ".extraction_done"
    if done.exists():
        log.info(t("reuse_extracted", d=workdir))
        return workdir
    log.info(t("extracting", z=inp.name, d=workdir))
    workdir.mkdir(parents=True, exist_ok=True)
    try:
        with zipfile.ZipFile(inp) as z:
            used: set[str] = set()
            for info in z.infolist():
                if info.is_dir():
                    continue
                name = Path(info.filename.replace("\\", "/")).name  # appiattisce, evita path traversal
                if not name or name.startswith("."):
                    continue
                base, ext, n = Path(name).stem, Path(name).suffix, 1
                while name.lower() in used:
                    n += 1
                    name = f"{base}_{n}{ext}"
                used.add(name.lower())
                with z.open(info) as src, open(workdir / name, "wb") as dst:
                    while chunk := src.read(1 << 20):
                        dst.write(chunk)
    except (zipfile.BadZipFile, EOFError) as e:
        raise UserError(t("err_zip", e=e))
    done.write_text("ok", encoding="utf-8")
    return workdir


def list_files(folder: Path) -> dict[str, Path]:
    files: dict[str, Path] = {}
    for p in sorted(folder.rglob("*")):
        if p.is_file() and not p.name.startswith("."):
            files.setdefault(p.name, p)
    return files


def chat_title(chat_file: Path, inp: Path) -> str:
    for src in (chat_file.stem, inp.stem):
        s = TITLE_PREFIX_RE.sub("", src).strip(" -_")
        if s and s.lower() not in ("_chat", "chat"):
            return s
    return "WhatsApp"


# --------------------------------------------------------------------------- #
# Elaborazione dei media
# --------------------------------------------------------------------------- #

def gui_group(kind: str) -> str:
    """Gruppi mostrati nell'interfaccia grafica."""
    return {"sticker": "image", "other": "document"}.get(kind, kind)


def process_media(names: list[str], files: dict[str, Path], args, cache: Cache) -> dict[str, dict]:
    ex = Extractor(args)
    results: dict[str, dict] = {}
    order = {"image": 0, "sticker": 0, "document": 1, "other": 1, "audio": 2, "video": 3}
    todo = []
    for n in names:
        kind = media_kind(n)
        if kind == "audio" and args.no_audio:
            results[n] = {"kind": kind, "meta": {}, "note": "n_audio_off", "size": files[n].stat().st_size}
            continue
        k = cache_key(args, n, files[n], kind)
        cached = cache.get(k)
        if cached is not None:
            results[n] = cached
        else:
            todo.append((order[kind], n, kind, k))
    todo.sort()
    log.info(t("media_summary", tot=len(names), cached=len(names) - len(todo), todo=len(todo)))

    # durata totale da trascrivere, per stimare il tempo residuo
    pending_av: dict[str, float] = {}
    for _, n, kind, _ in todo:
        if kind == "audio" or (kind == "video" and not args.no_video):
            try:
                pending_av[n] = probe_av(files[n]).get("duration") or 0
            except Exception:
                pending_av[n] = 0
    tot_av = sum(pending_av.values())
    if tot_av:
        log.info(t("audio_total", d=fmt_dur(tot_av), n=len(pending_av)))

    # piano di lavoro per l'interfaccia grafica: totali per gruppo e quanti già pronti (cache)
    totals, ready = Counter(), Counter()
    for n in names:
        g = gui_group(media_kind(n))
        totals[g] += 1
        if n in results:
            ready[g] += 1
    emit("plan", totals=dict(totals), ready=dict(ready), todo=len(todo), av_total=tot_av)

    # Stima del tempo: ogni audio costa un tempo fisso (rilevamento lingua, avvio) più una parte
    # proporzionale alla durata. I due valori si correggono con le misure reali man mano che si procede.
    base_fixed, base_rate = 6.0, 0.45
    fixed, rate, other_avg, n_other, spent_other = base_fixed, base_rate, 0.3, 0, 0.0
    av_samples: list[tuple[float, float]] = []  # (durata, secondi impiegati)

    def refit() -> tuple[float, float]:
        if len(av_samples) >= 4:
            n_s = len(av_samples)
            mx = sum(d for d, _ in av_samples) / n_s
            my = sum(e for _, e in av_samples) / n_s
            var = sum((d - mx) ** 2 for d, _ in av_samples)
            if var > 25:
                r = sum((d - mx) * (e - my) for d, e in av_samples) / var
                f = my - r * mx
                if 0.05 <= r <= 5 and 0 <= f <= 60:
                    return f, r
        pred = sum(base_fixed + base_rate * d for d, _ in av_samples)
        k = sum(e for _, e in av_samples) / pred if pred > 0 else 1.0
        return base_fixed * k, base_rate * k

    def remaining(from_idx: int) -> float:
        rest = todo[from_idx:]
        return sum(fixed + pending_av[n] * rate if n in pending_av else other_avg for _, n, _, _ in rest)

    done_av, n_av_done, t_start = 0.0, 0, time.time()
    for i, (_, n, kind, k) in enumerate(todo, 1):
        path = files[n]
        emit("file_start", i=i, n=len(todo), kind=gui_group(kind), name=n, duration=pending_av.get(n, 0),
             eta=remaining(i - 1))
        t0 = time.time()
        try:
            res = ex.run(kind, path)
            res["kind"] = kind
            res["size"] = path.stat().st_size
            cache.put(k, res)  # gli errori non vanno in cache: al prossimo avvio si riprova
            status = t("status_ok")
        except KeyboardInterrupt:
            raise
        except Exception as e:
            res = {"kind": kind, "meta": {}, "size": path.stat().st_size,
                   "error": f"{type(e).__name__}: {e}".replace(str(path), n).replace(repr(str(path))[1:-1], n)[:300]}
            log.debug(f"{n}\n{traceback.format_exc()}")
            status = f"{t('status_err')}: {res['error']}"
        results[n] = res
        el = time.time() - t0
        extra = ""
        if n in pending_av:
            done_av += pending_av[n]
            # il primo file include il caricamento del modello: non va usato per la stima
            if n_av_done > 0 or len(pending_av) == 1:
                av_samples.append((pending_av[n], el))
                fixed, rate = refit()
            n_av_done += 1
            if done_av > 0 and tot_av > done_av:
                extra = t("eta", a=fmt_dur(done_av), b=fmt_dur(tot_av), c=fmt_dur(remaining(i)))
        else:
            n_other += 1
            spent_other += el
            other_avg = spent_other / n_other
        log.info(f"[{i}/{len(todo)}] {kind:8s} {n} ({el:.1f}s) {status}{extra}")
        emit("file_done", i=i, n=len(todo), kind=gui_group(kind), name=n, ok="error" not in res,
             error=res.get("error"), eta=remaining(i), elapsed=time.time() - t_start)
    if todo:
        log.info(t("media_done", d=fmt_dur(time.time() - t_start)))
    return results


# --------------------------------------------------------------------------- #
# Generazione del Markdown
# --------------------------------------------------------------------------- #

class Renderer:
    def __init__(self, results: dict[str, dict], args):
        self.results, self.args = results, args
        self.appendix: list[tuple[str, str, str]] = []  # (id, nome file, testo)

    def media_block(self, name: str) -> list[str]:
        res = self.results.get(name)
        if res is None:
            return [f"[{t('k_missing')}] `{name}`"]
        kind = res.get("kind", media_kind(name))
        label = t("k_" + kind)
        if kind == "audio" and re.match(r"(PTT-|.*-AUDIO-)", name, re.I):
            label = t("k_voice")
        meta = res.get("meta", {})
        info = [f"[{label}] `{name}`"]
        if meta.get("duration"):
            info.append(t("m_duration", v=fmt_dur(meta["duration"])))
        for key in ("resolution", "dimensions"):
            if key in meta:
                info.append(meta[key])
        if meta.get("animated"):
            info.append(t("m_animated"))
        if meta.get("language"):
            info.append(t("m_language", v=meta["language"]))
        for key in ("pages", "slides", "sheets", "ocr_pages"):
            if key in meta:
                info.append(t("m_" + key, v=meta[key]))
        if kind in ("document", "other"):
            info.append(fmt_size(res.get("size", 0)))
        for key in ("title", "author", "created", "taken", "camera", "gps", "recorded"):
            if meta.get(key):
                info.append(t("m_" + key, v=meta[key]))
        out = [" · ".join(info)]

        if res.get("error"):
            out.append("> " + t("x_error", e=res["error"]))
            return out
        if res.get("note"):
            out.append(f"> ({t(res['note'])})")
        segs = res.get("segments")
        if segs is not None:
            if not segs:
                out.append("> " + t("x_no_speech"))
            elif (meta.get("duration") or 0) > 90:
                out.append(f"> {t('x_transcript')}:")
                out += [f"> [{fmt_dur(s)}] {txt}" for s, txt in segs]
            else:
                out += quote(f"{t('x_transcript')}: " + " ".join(txt for _, txt in segs))
        text = (res.get("text") or "").strip()
        if text:
            head = t("x_ocr") if kind in ("image", "sticker") else t("x_content")
            limit = self.args.max_doc_chars
            if len(text) > self.args.appendix_threshold:
                aid = f"D{len(self.appendix) + 1}"
                note = t("x_truncated", a=limit, b=len(text)) if len(text) > limit else ""
                self.appendix.append((aid, name, text[:limit] + ("\n[...]" if note else "")))
                excerpt = text[:1200].rsplit(" ", 1)[0]
                out += quote(t("x_excerpt", h=head, id=aid, n=note) + "\n" + excerpt + " [...]")
            else:
                out += quote(f"{head}:\n{text}")
        return out

    def message(self, m: Message) -> list[str]:
        text_lines = [md_escape_line(l) for l in m.text.split("\n")] if m.text else []
        if m.sender is None and not m.attachments:
            return [f"*[{m.ts:%H:%M}] {t('system')} {' '.join(text_lines)}*"]
        prefix = f"**[{m.ts:%H:%M}] {m.sender or t('system')}:**"
        lines: list[str] = []
        if m.attachments:
            for i, a in enumerate(m.attachments):
                blk = self.media_block(a)
                lines.append(f"{prefix} {blk[0]}" if i == 0 else blk[0])
                if i == 0 and text_lines:  # didascalia subito sotto l'allegato, prima del contenuto estratto
                    lines.append(t("x_caption") + text_lines[0])
                    lines += text_lines[1:]
                lines += blk[1:]
        else:
            lines.append(f"{prefix} {text_lines[0] if text_lines else ''}".rstrip())
            lines += text_lines[1:]
        return lines


def build_markdown(msgs: list[Message], title: str, results: dict, unreferenced: list[str],
                   missing: list[str], args, source: Path, part_label: str | None = None) -> str:
    r = Renderer(results, args)
    L: list[str] = []
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    L += ["# " + t("md_title", t=title) + (f" — {part_label}" if part_label else ""), ""]
    L += [t("md_generated", v=VERSION, d=now, s=source.name), ""]

    # ---- riepilogo ----------------------------------------------------------
    L += ["## " + t("md_summary"), ""]
    if msgs:
        first, last = msgs[0].ts, msgs[-1].ts
        L.append("- " + t("md_period", a=f"{first:%Y-%m-%d %H:%M}", b=f"{last:%Y-%m-%d %H:%M}",
                              days=(last.date() - first.date()).days + 1, active=len({m.ts.date() for m in msgs})))
    L.append("- " + t("md_messages", n=len(msgs), s=sum(1 for m in msgs if m.sender is None)))
    per: dict[str, list[Message]] = defaultdict(list)
    for m in msgs:
        if m.sender:
            per[m.sender].append(m)
    if per:
        L.append("- " + t("md_participants"))
        for s, ms in sorted(per.items(), key=lambda x: -len(x[1])):
            L.append("  - " + t("md_participant", s=s, n=len(ms), w=sum(len(m.text.split()) for m in ms),
                                  a=sum(len(m.attachments) for m in ms),
                                  f=f"{ms[0].ts:%Y-%m-%d}", l=f"{ms[-1].ts:%Y-%m-%d}"))
    refd = [a for m in msgs for a in m.attachments]
    if refd:
        kinds = Counter(results[a]["kind"] for a in refd if a in results)
        parts = [f"{t('k_' + k).lower()}: {v}" for k, v in kinds.most_common()]
        n_missing = sum(1 for a in refd if a not in results)
        if n_missing:
            parts.append(t("md_not_included", n=n_missing))
        aud = sum((results[a].get("meta", {}).get("duration") or 0) for a in refd
                  if a in results and results[a].get("kind") == "audio")
        L.append("- " + t("md_attachments", n=len(refd), p=", ".join(parts))
                 + (t("md_audio_total", d=fmt_dur(aud)) if aud else ""))
    n_omitted, n_omitted_voice = count_omitted(msgs)
    if n_omitted:
        L.append("- " + t("md_omitted", n=n_omitted, v=n_omitted_voice))
    errs = [a for a in set(refd) | set(unreferenced) if a in results and results[a].get("error")]
    if errs:
        L.append("- " + t("md_unreadable", n=len(errs)))
    years = Counter(m.ts.year for m in msgs)
    if len(years) > 1:
        L.append("- " + t("md_per_year") + ", ".join(f"{y}: {c}" for y, c in sorted(years.items())))
    L.append("")

    # ---- legenda per il lettore (umano o AI) ---------------------------------
    L += ["## " + t("md_howto"), ""]
    L += ["- " + x for x in t("md_howto_lines")]
    L.append("")

    # ---- conversazione --------------------------------------------------------
    L.append("## " + t("md_conversation"))
    days = t("weekdays")
    cur_day = None
    for m in msgs:
        day = m.ts.date()
        if day != cur_day:
            cur_day = day
            L += ["", f"### {day.isoformat()}, {days[day.weekday()]}"]
        L.append("")
        L += r.message(m)
    L.append("")

    # ---- media non citati ------------------------------------------------------
    if unreferenced:
        L += ["## " + t("md_unref"), "", t("md_unref_note")]
        for n in unreferenced:
            blk = r.media_block(n)
            d = date_from_filename(n)
            L += ["", f"- ({d}) {blk[0]}" if d else f"- {blk[0]}"]
            L += ["  " + x for x in blk[1:]]
        L.append("")
    if missing:
        L += ["## " + t("md_missing"), "", ", ".join(f"`{n}`" for n in missing), ""]

    # ---- appendice ------------------------------------------------------------
    if r.appendix:
        L.append("## " + t("md_appendix"))
        for aid, name, text in r.appendix:
            L += ["", f"### {t('md_appendix_item', id=aid)}: `{name}`", "", "```text",
                  text.replace("```", "ˋˋˋ"), "```"]
        L.append("")
    return "\n".join(L)


def write_atomic(path: Path, text: str) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8", newline="\n")
    os.replace(tmp, path)


# --------------------------------------------------------------------------- #
# Interazione e controlli
# --------------------------------------------------------------------------- #

def system_ui_language() -> str:
    """Lingua dell'interfaccia di Windows, se è una delle tre supportate (altrimenti inglese)."""
    try:
        import ctypes

        primary = ctypes.windll.kernel32.GetUserDefaultUILanguage() & 0x3FF
        return {0x10: "it", 0x16: "pt", 0x09: "en"}.get(primary, "en")
    except Exception:
        return "en"


def ask_ui_language() -> str:
    default = system_ui_language()
    codes = ["it", "pt", "en"]
    print()
    print("  zapp2md " + VERSION)
    print()
    for i, c in enumerate(codes, 1):
        print(f"  {i}) {LANGS[c]}" + ("  *" if c == default else ""))
    while True:
        try:
            c = input(f"  Lingua / Idioma / Language [1-3, Enter = {codes.index(default) + 1}]: ").strip().lower()
        except EOFError:
            return default
        m = {"1": "it", "2": "pt", "3": "en", "it": "it", "pt": "pt", "en": "en", "": default}
        if c in m:
            print()
            return m[c]


def ask_input_path() -> str | None:
    print(t("ask_file"))
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        p = filedialog.askopenfilename(title="WhatsApp .zip",
                                       filetypes=[("WhatsApp export", "*.zip *.txt"), ("*", "*.*")])
        root.destroy()
        if p:
            return p
    except Exception:
        pass
    try:
        p = input(t("ask_path")).strip().strip('"').strip("'").strip()
    except EOFError:
        return None
    return p or None


def ask_yes(prompt: str, default: bool = True) -> bool:
    try:
        a = input(prompt).strip().lower()
    except EOFError:
        return default
    return default if not a else a[0] in t("yes_letters")


def setup_check(args) -> int:
    """Controllo dell'installazione e download anticipato del modello (usato da install.bat)."""
    ok = True
    checks = [
        ("PyAV", lambda: __import__("av")),
        ("pypdfium2", lambda: __import__("pypdfium2")),
        ("python-docx / openpyxl / python-pptx", lambda: (__import__("docx"), __import__("openpyxl"), __import__("pptx"))),
        (f"Whisper '{args.model}'", lambda: Transcriber(args.model, args.speech_lang, args.threads)),
        ("Windows OCR", lambda: print("   OCR:", ", ".join(WinOCR(args.speech_lang).available))),
    ]
    for name, fn in checks:
        try:
            fn()
            print(f"   OK  {name}")
        except Exception as e:
            print("   " + t("setup_fail", c=name, e=e))
            if "OCR" in name:
                # senza OCR il programma funziona comunque (salta solo il testo nelle immagini)
                print("   " + t("ocr_off", e=e))
                continue
            ok = False
            if "DLL" in str(e):
                print("   " + t("setup_longpath"))
    print()
    print(t("setup_ok") if ok else "!!!")
    return 0 if ok else 1


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def main(argv=None) -> int:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(
        description="WhatsApp export (zip/folder) -> Markdown for LLMs (Claude, ChatGPT, Gemini...). Italiano / Português / English.")
    ap.add_argument("input", nargs="?", help=".zip exported from WhatsApp, the extracted folder, or the chat .txt")
    ap.add_argument("-o", "--output", help="output .md file (default: next to the input)")
    ap.add_argument("--ui", choices=list(LANGS), help="interface and document language: it, pt, en")
    ap.add_argument("--speech-lang", help="main spoken language of the chat (it, pt, en, es, fr, de, ... or auto); "
                                          "default = --ui language")
    ap.add_argument("--model", default=DEFAULT_MODEL,
                    help="Whisper model: large-v3-turbo (default), large-v3 (best, slow), medium, small (fast)")
    ap.add_argument("--no-video", action="store_true", help="do not transcribe video audio")
    ap.add_argument("--video", action="store_true", help="transcribe video audio without asking")
    ap.add_argument("--no-audio", action="store_true", help="do not transcribe audio (quick test)")
    ap.add_argument("--no-ocr", action="store_true", help="do not read text in images")
    ap.add_argument("--split-by-year", action="store_true", help="also write one .md file per year")
    ap.add_argument("--date-format", default="auto", choices=["auto", "dmy", "mdy", "ymd"],
                    help="date order in the chat file (default: automatic)")
    ap.add_argument("--max-ocr-pages", type=int, default=50, help="max scanned pages OCR'd per PDF")
    ap.add_argument("--max-doc-chars", type=int, default=100_000, help="max characters kept per document")
    ap.add_argument("--appendix-threshold", type=int, default=3000,
                    help="documents longer than this go to the appendix")
    ap.add_argument("--threads", type=int, default=max(1, os.cpu_count() or 4), help="CPU threads for transcription")
    ap.add_argument("--setup", action="store_true", help="check installation and download the model")
    ap.add_argument("--progress-json", action="store_true", help=argparse.SUPPRESS)  # usato dall'interfaccia grafica
    ap.add_argument("--check-model", action="store_true", help=argparse.SUPPRESS)  # usato da setup.ps1
    args = ap.parse_args(argv)

    if args.check_model:
        # Il modello è già sul PC ed è integro? (nessun download: solo file locali)
        try:
            from faster_whisper import WhisperModel

            WhisperModel(args.model, device="cpu", compute_type="int8", local_files_only=True)
            return 0
        except Exception as e:
            print(f"model not available locally: {e}")
            return 1

    global _JSON_EVENTS
    _JSON_EVENTS = args.progress_json
    interactive = (sys.stdin is not None and sys.stdin.isatty()) and not args.progress_json
    asked = False
    if args.ui is None:
        if interactive and not args.setup:
            args.ui = ask_ui_language()
            asked = True
        else:
            args.ui = system_ui_language()
    set_lang(args.ui)
    args.speech_lang = (args.speech_lang or args.ui).lower()

    if args.setup:
        return setup_check(args)

    if not args.input:
        if not interactive:
            ap.print_usage()
            return 2
        args.input = ask_input_path()
        asked = True
        if not args.input:
            return 2
    if asked and not args.no_video and not args.video:
        args.no_video = not ask_yes(t("ask_video"))

    try:
        return run(args)
    except UserError as e:
        if _JSON_EVENTS:
            emit("fatal", msg=str(e))
        else:
            print(str(e))
        return 2
    except Exception as e:
        if _JSON_EVENTS:
            emit("fatal", msg=f"{type(e).__name__}: {e}", trace=traceback.format_exc())
            return 1
        raise


def run(args) -> int:
    inp = Path(args.input.strip().strip('"')).expanduser().resolve()
    if not inp.exists():
        raise UserError(t("err_notfound", p=inp))
    base = inp.with_suffix("") if inp.is_file() else inp
    out_md = Path(args.output).resolve() if args.output else base.parent / (base.name + ".md")
    out_md.parent.mkdir(parents=True, exist_ok=True)
    log_path = out_md.with_suffix(".log.txt")

    log.setLevel(logging.DEBUG)
    log.handlers.clear()
    fh = logging.FileHandler(log_path, encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    ch = JsonLogHandler() if _JSON_EVENTS else logging.StreamHandler(sys.stdout)
    ch.setFormatter(logging.Formatter("%(asctime)s  %(message)s", "%H:%M:%S"))
    ch.setLevel(logging.INFO)
    log.addHandler(fh)
    log.addHandler(ch)
    for noisy in ("faster_whisper", "huggingface_hub", "httpx"):
        logging.getLogger(noisy).setLevel(logging.WARNING)

    log.info(t("start", v=VERSION, p=inp))
    workdir = out_md.parent / (out_md.stem + "_media")
    folder = prepare_input(inp, workdir)
    files = list_files(folder)
    chat_file = inp if inp.suffix.lower() == ".txt" else find_chat_file(files)
    files.pop(chat_file.name, None)
    title = chat_title(chat_file, inp)
    log.info(t("chat_found", t=title, f=chat_file.name, n=len(files)))

    msgs, order = parse_chat(chat_file, files, args.date_format)
    if not msgs:
        log.error(t("no_messages"))
        emit("fatal", msg=t("no_messages"))
        return 3
    unsorted = sum(1 for a, b in zip(msgs, msgs[1:]) if b.ts < a.ts)
    msgs.sort(key=lambda m: m.ts)  # ordinamento stabile: a parità di orario resta l'ordine originale
    log.info(t("messages_found", n=len(msgs), o=order, a=f"{msgs[0].ts:%Y-%m-%d}", b=f"{msgs[-1].ts:%Y-%m-%d}",
               x=t("reordered", n=unsorted) if unsorted else ""))

    referenced: list[str] = []
    seen: set[str] = set()
    for m in msgs:
        for a in m.attachments:
            if a not in seen:
                seen.add(a)
                referenced.append(a)
    missing = [a for a in referenced if a not in files]
    unreferenced = sorted(n for n in files if n not in seen)
    to_process = [a for a in referenced if a in files] + unreferenced

    cache = Cache(out_md.parent / (out_md.stem + ".cache.sqlite"))
    try:
        results = process_media(to_process, files, args, cache)
    except KeyboardInterrupt:
        log.warning(t("interrupted"))
        return 130
    finally:
        cache.close()

    write_atomic(out_md, build_markdown(msgs, title, results, unreferenced, missing, args, inp))
    written = [out_md]
    if args.split_by_year:
        for y in sorted({m.ts.year for m in msgs}):
            p = out_md.with_name(f"{out_md.stem}_{y}.md")
            write_atomic(p, build_markdown([m for m in msgs if m.ts.year == y], title, results, [], [],
                                           args, inp, part_label=str(y)))
            written.append(p)

    n_err = sum(1 for r in results.values() if r.get("error"))
    n_omitted, n_omitted_voice = count_omitted(msgs)
    log.info("-" * 70)
    no_media = n_omitted > 0 and not files  # nessun file allegato: esportata con "Senza media"
    if no_media:
        log.warning(t("warn_omitted", n=n_omitted, v=n_omitted_voice))
        log.info("-" * 70)
    elif n_omitted:
        log.warning(t("warn_omitted_partial", n=n_omitted,
                      v=t("voice_part", v=n_omitted_voice) if n_omitted_voice else ""))
    for p in written:
        size = p.stat().st_size
        log.info(t("created", p=p, s=fmt_size(size), tok=f"{size // 4:,}".replace(",", ".")))
    if out_md.stat().st_size // 4 > 400_000 and not args.split_by_year:
        log.info(t("hint_split"))
    if n_err:
        log.info(t("n_errors", n=n_err, f=log_path.name))
    log.info(t("done"))
    emit("result", outputs=[str(p) for p in written], log=str(log_path),
         errors=[[n, r["error"]] for n, r in results.items() if r.get("error")],
         omitted=n_omitted, omitted_voice=n_omitted_voice, no_media=no_media, tokens=out_md.stat().st_size // 4,
         messages=len(msgs))
    return 0


if __name__ == "__main__":
    code = main()
    sys.exit(code)
